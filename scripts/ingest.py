#!/usr/bin/env python3
"""Bounded text extraction for common Source to Motion inputs.

This script extracts source material; it never writes marketing claims.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import os
import re
import shutil
import ssl
import subprocess
import tempfile
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import quote, urlparse
from urllib.request import Request, urlopen

import certifi

MAX_BYTES = 8_000_000
MAX_TEXT = 180_000
USER_AGENT = "source-to-motion/0.1 (+https://github.com/Balance0014/source-to-motion)"


class PageText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "svg", "nav", "footer", "aside", "form"}:
            self.hidden += 1
        if not self.hidden and tag in {"p", "h1", "h2", "h3", "h4", "li", "tr", "br", "section"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "svg", "nav", "footer", "aside", "form"} and self.hidden:
            self.hidden -= 1

    def handle_data(self, data):
        if not self.hidden:
            value = re.sub(r"\s+", " ", data).strip()
            if value:
                self.parts.append(" " + value)


def fetch(url: str, token: str | None = None) -> bytes:
    headers = {"User-Agent": USER_AGENT}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    request = Request(url, headers=headers)
    with urlopen(request, timeout=15, context=ssl.create_default_context(cafile=certifi.where())) as response:
        length = response.headers.get("Content-Length")
        if length and int(length) > MAX_BYTES:
            raise ValueError("Source is too large for the bounded extractor")
        data = response.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ValueError("Source exceeds the 8 MB extraction limit")
    return data


def github_readme(url: str) -> tuple[str, str]:
    path = urlparse(url).path.strip("/").split("/")
    if len(path) < 2:
        raise ValueError("GitHub URL needs owner/repository")
    owner, repo = path[:2]
    if len(path) >= 5 and path[2] == "blob":
        raw = f"https://raw.githubusercontent.com/{owner}/{repo}/{path[3]}/{'/'.join(path[4:])}"
        return fetch(raw).decode("utf-8", "replace"), raw
    token = os.getenv("GITHUB_TOKEN") or os.getenv("GH_TOKEN")
    metadata = json.loads(fetch(f"https://api.github.com/repos/{owner}/{repo}", token))
    branch = metadata["default_branch"]
    commit = json.loads(fetch(f"https://api.github.com/repos/{owner}/{repo}/commits/{quote(branch, safe='')}", token))["sha"]
    for name in ("README.md", "readme.md", "README.rst", "README.txt"):
        raw = f"https://raw.githubusercontent.com/{owner}/{repo}/{commit}/{name}"
        try:
            return fetch(raw).decode("utf-8", "replace"), raw
        except Exception:
            continue
    raise ValueError("Repository README was not found; inspect the repository with your agent")


def extract_pdf(path: Path) -> tuple[str, bool]:
    from pypdf import PdfReader

    pages = PdfReader(str(path)).pages
    if len(pages) > 60:
        raise ValueError("PDF exceeds 60 pages; select the relevant pages first")
    chunks = [f"\n[Page {i+1}]\n{page.extract_text() or ''}" for i, page in enumerate(pages)]
    text = "\n".join(chunks)
    if len(text.strip()) >= 80:
        return text, False
    if not (shutil.which("pdftoppm") and shutil.which("tesseract")):
        raise ValueError("Image-only PDF: OCR is required. Use agent vision or install pdftoppm + tesseract.")
    with tempfile.TemporaryDirectory(prefix="stm-ocr-") as tmp:
        prefix = str(Path(tmp) / "page")
        subprocess.run(["pdftoppm", "-f", "1", "-l", str(min(10, len(pages))), "-r", "160", "-png", str(path), prefix],
                       check=True, timeout=90, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        ocr = []
        for i, image in enumerate(sorted(Path(tmp).glob("page-*.png"))):
            result = subprocess.run(["tesseract", str(image), "stdout"], capture_output=True, text=True, timeout=25, check=True)
            ocr.append(f"\n[OCR Page {i+1}]\n{result.stdout}")
    text = "\n".join(ocr)
    if len(text.strip()) < 30:
        raise ValueError("OCR returned too little text; inspect the PDF visually")
    return text, True


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table_index, table in enumerate(doc.tables, 1):
        parts.append(f"[Table {table_index}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.replace("\n", " ") for cell in row.cells))
    return "\n".join(parts)


def extract_image(path: Path) -> tuple[str, bool]:
    if not shutil.which("tesseract"):
        raise ValueError("Image text needs OCR. Use agent vision or install tesseract.")
    result = subprocess.run(["tesseract", str(path), "stdout"], capture_output=True, text=True, timeout=25, check=True)
    if len(result.stdout.strip()) < 20:
        raise ValueError("OCR returned too little text; inspect the image visually")
    return result.stdout, True


def extract(source: str) -> dict:
    parsed = urlparse(source)
    is_url = parsed.scheme in {"http", "https"}
    ocr = False
    if is_url and parsed.hostname == "github.com" and len(parsed.path.strip("/").split("/")) >= 2:
        text, locator = github_readme(source)
        kind = "github-file" if "/blob/" in parsed.path else "github-readme"
    elif is_url:
        data = fetch(source)
        suffix = Path(parsed.path).suffix.lower()
        if suffix == ".pdf":
            with tempfile.TemporaryDirectory(prefix="stm-pdf-") as tmp:
                local = Path(tmp) / "source.pdf"
                local.write_bytes(data)
                text, ocr = extract_pdf(local)
            kind = "pdf"
        elif suffix == ".docx":
            with tempfile.TemporaryDirectory(prefix="stm-docx-") as tmp:
                local = Path(tmp) / "source.docx"
                local.write_bytes(data)
                text = extract_docx(local)
            kind = "docx"
        elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            with tempfile.TemporaryDirectory(prefix="stm-image-") as tmp:
                local = Path(tmp) / f"source{suffix}"
                local.write_bytes(data)
                text, ocr = extract_image(local)
            kind = "image"
        elif suffix in {".txt", ".md", ".rst", ".csv", ".json"}:
            text = data.decode("utf-8", "replace")
            kind = "text"
        else:
            parser = PageText()
            parser.feed(data.decode("utf-8", "replace"))
            text = "".join(parser.parts)
            kind = "website"
        locator = source
    else:
        path = Path(source).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(path)
        if path.stat().st_size > MAX_BYTES:
            raise ValueError("File exceeds the 8 MB extraction limit")
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text, ocr = extract_pdf(path)
            kind = "pdf"
        elif suffix == ".docx":
            text = extract_docx(path)
            kind = "docx"
        elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            text, ocr = extract_image(path)
            kind = "image"
        elif suffix in {".txt", ".md", ".rst", ".csv", ".json"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            kind = "text"
        else:
            raise ValueError(f"Unsupported file type: {suffix or '(none)'}")
        locator = str(path)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 30:
        raise ValueError("The source yielded too little text; inspect it with agent vision or provide another source")
    if len(text) > MAX_TEXT:
        text = text[:MAX_TEXT] + "\n[TRUNCATED — select relevant sections]"
    return {"source": source, "locator": locator, "kind": kind, "ocr": ocr,
            "retrieved_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"), "text": text}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("source", help="URL or local PDF/DOCX/TXT/MD/image file")
    p.add_argument("--out", required=True, help="Output JSON path")
    args = p.parse_args()
    result = extract(args.source)
    output = Path(args.out)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k != "text"} | {"characters": len(result["text"])}))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        raise SystemExit(f"ingest failed: {exc}")
